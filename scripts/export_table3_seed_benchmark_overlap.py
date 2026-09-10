import csv
import json
from collections import OrderedDict
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parents[1]
DATA_DIR = ROOT / "data_set"
OUTPUT_DIR = ROOT / "output" / "report_exports" / "20260910_table3_seed_benchmark_overlap"

RUNS = {
    "BRCA": {
        "stamp": "20260905_194936",
        "expected_alpha": 0.9667,
        "expected_beta": 0.2396,
        "seed_path": DATA_DIR / "seed_set" / "New" / "TCGA-BRCA_seed.txt",
        "benchmark_path": DATA_DIR / "Onco_KB BRCA.csv",
    },
    "COAD": {
        "stamp": "20260905_195534",
        "expected_alpha": 0.7465,
        "expected_beta": 0.2123,
        "seed_path": DATA_DIR / "seed_set" / "New" / "TCGA-COAD_seed.txt",
        "benchmark_path": DATA_DIR / "Onco_KB_COAD.csv",
    },
    "LUAD": {
        "stamp": "20260905_200214",
        "expected_alpha": 0.7030,
        "expected_beta": 0.3636,
        "seed_path": DATA_DIR / "seed_set" / "New" / "TCGA-LUAD_seed.txt",
        "benchmark_path": DATA_DIR / "Onco_KB_LUAD.csv",
    },
    "THCA": {
        "stamp": "20260905_200835",
        "expected_alpha": 0.8872,
        "expected_beta": 0.1705,
        "seed_path": DATA_DIR / "seed_set" / "New" / "TCGA-THCA_seed.txt",
        "benchmark_path": DATA_DIR / "Onco_KB_THCA.csv",
    },
}


def read_lines(path):
    values = []
    with open(path, encoding="utf-8-sig") as handle:
        for line in handle:
            value = line.strip()
            if value:
                values.append(value)
    return values


def read_gene_symbols_csv(path):
    genes = []
    with open(path, encoding="utf-8-sig", newline="") as handle:
        reader = csv.DictReader(handle)
        for row in reader:
            gene = (row.get("Gene") or "").strip().upper()
            if gene:
                genes.append(gene)
    return list(OrderedDict.fromkeys(genes))


def load_ensembl_mapping(path):
    mapping = {}
    with open(path, encoding="utf-8-sig", newline="") as handle:
        reader = csv.DictReader(handle, delimiter="\t")
        for row in reader:
            ensembl = (row.get("Gene stable ID") or "").strip().split(".")[0]
            symbol = (row.get("Gene name") or "").strip().upper()
            if ensembl and symbol and ensembl not in mapping:
                mapping[ensembl] = symbol
    return mapping


def normalize_identifier(value):
    return (value or "").strip().split(".")[0]


def map_symbol(identifier, mapping):
    base = normalize_identifier(identifier)
    return mapping.get(base, base.upper() if not base.startswith("ENSG") else "")


def run_dir(disease, stamp):
    return ROOT / "output" / disease / "optuna_biorank_compare" / "full" / stamp


def ranking_path_from_summary(disease, stamp, expected_alpha, expected_beta):
    summary_path = run_dir(disease, stamp) / "optimization_summary.json"
    with open(summary_path, encoding="utf-8") as handle:
        summary = json.load(handle)
    if summary.get("validation_mode") != "Disease OncoKB":
        raise ValueError(f"{disease} run does not use Disease OncoKB: {summary_path}")
    for candidate in summary.get("selected_candidates", []):
        if int(candidate.get("candidate_rank", 0)) != 1:
            continue
        alpha = round(float(candidate["alpha"]), 4)
        beta = round(float(candidate["beta"]), 4)
        if alpha != expected_alpha or beta != expected_beta:
            raise ValueError(
                f"{disease} candidate #1 alpha/beta mismatch: got {alpha},{beta}; "
                f"expected {expected_alpha},{expected_beta}"
            )
        return run_dir(disease, stamp) / "rankings" / Path(candidate["ranking_path"]).name, summary_path
    raise ValueError(f"{disease} candidate #1 not found in {summary_path}")


def read_ranking(path, mapping):
    rows = []
    with open(path, encoding="utf-8-sig", newline="") as handle:
        reader = csv.DictReader(handle, delimiter="\t")
        for rank, row in enumerate(reader, 1):
            identifier = (row.get("GeneNames") or row.get("name") or "").strip()
            symbol = map_symbol(identifier, mapping)
            score_text = row.get("Score") or row.get("score") or ""
            try:
                score = float(score_text)
            except ValueError:
                score = ""
            rows.append(
                {
                    "Rank": rank,
                    "GeneIdentifier": identifier,
                    "GeneSymbol": symbol,
                    "Score": score,
                }
            )
    return rows


def write_tsv(path, rows, headers):
    with open(path, "w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=headers, delimiter="\t")
        writer.writeheader()
        writer.writerows(rows)


def rel(path):
    try:
        return str(Path(path).relative_to(ROOT))
    except ValueError:
        return str(path)


def style_sheet(sheet):
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    thin = Side(style="thin", color="D9E2F3")
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in sheet.iter_rows():
        for cell in row:
            cell.border = Border(top=thin, bottom=thin, left=thin, right=thin)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if isinstance(cell.value, float):
                cell.number_format = "0.00"
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for index, column in enumerate(sheet.columns, 1):
        max_len = max(len(str(cell.value)) if cell.value is not None else 0 for cell in column)
        sheet.column_dimensions[get_column_letter(index)].width = min(max(max_len + 2, 12), 56)


def main():
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    details_dir = OUTPUT_DIR / "per_disease_details"
    details_dir.mkdir(exist_ok=True)
    mapping = load_ensembl_mapping(DATA_DIR / "mart_biotool.txt")

    table_rows = []
    manifest_rows = []
    seed_detail_rows = []
    top100_detail_rows = []
    benchmark_detail_rows = []

    for disease, config in RUNS.items():
        ranking_path, summary_path = ranking_path_from_summary(
            disease,
            config["stamp"],
            config["expected_alpha"],
            config["expected_beta"],
        )
        seed_identifiers = list(OrderedDict.fromkeys(read_lines(config["seed_path"])))
        seed_identifier_set = {normalize_identifier(value) for value in seed_identifiers}
        seed_rows = []
        for identifier in seed_identifiers:
            base = normalize_identifier(identifier)
            symbol = map_symbol(identifier, mapping)
            seed_rows.append(
                {
                    "Cancer": disease,
                    "SeedIdentifier": identifier,
                    "SeedIdentifierBase": base,
                    "GeneSymbol": symbol,
                }
            )
        seed_symbols = {row["GeneSymbol"] for row in seed_rows if row["GeneSymbol"]}
        benchmark_symbols = set(read_gene_symbols_csv(config["benchmark_path"]))
        ranking_rows = read_ranking(ranking_path, mapping)

        overlap = seed_symbols & benchmark_symbols
        benchmark_without_seeds = benchmark_symbols - seed_symbols

        nonseed_ranking = []
        for row in ranking_rows:
            base = normalize_identifier(row["GeneIdentifier"])
            symbol = row["GeneSymbol"]
            is_seed = base in seed_identifier_set or symbol in seed_symbols
            if is_seed:
                continue
            nonseed_ranking.append(
                {
                    "Cancer": disease,
                    "NonSeedRank": len(nonseed_ranking) + 1,
                    "OriginalRank": row["Rank"],
                    "GeneIdentifier": row["GeneIdentifier"],
                    "GeneSymbol": symbol,
                    "Score": row["Score"],
                    "InBenchmarkExcludingSeeds": "Yes" if symbol in benchmark_without_seeds else "No",
                }
            )
            if len(nonseed_ranking) == 100:
                break

        top100_symbols = {row["GeneSymbol"] for row in nonseed_ranking if row["GeneSymbol"]}
        exclusion_matches = top100_symbols & benchmark_without_seeds
        recall_excluding = (
            len(exclusion_matches) / len(benchmark_without_seeds)
            if benchmark_without_seeds
            else "NA"
        )
        seed_overlap_pct = len(overlap) / len(seed_symbols) * 100 if seed_symbols else 0.0
        benchmark_coverage_pct = len(overlap) / len(benchmark_symbols) * 100 if benchmark_symbols else 0.0

        table_rows.append(
            {
                "Cancer": disease,
                "Enriched seeds": len(seed_symbols),
                "|Omega_c|": len(benchmark_symbols),
                "|S_c intersect Omega_c|": len(overlap),
                "Seed overlap (%)": seed_overlap_pct,
                "Benchmark coverage (%)": benchmark_coverage_pct,
                "Matches@100 excluding seeds": len(exclusion_matches),
                "Recall@100 excluding seeds": recall_excluding,
            }
        )
        manifest_rows.append(
            {
                "Cancer": disease,
                "Seed file": rel(config["seed_path"]),
                "Benchmark OncoKB file": rel(config["benchmark_path"]),
                "Ranking file": rel(ranking_path),
                "Optimization summary": rel(summary_path),
                "Alpha": config["expected_alpha"],
                "Beta": config["expected_beta"],
                "Seed identifiers": len(seed_identifiers),
                "Mapped seed symbols": len(seed_symbols),
                "Benchmark genes": len(benchmark_symbols),
            }
        )
        for row in seed_rows:
            enriched = dict(row)
            enriched["InCancerSpecificOncoKB"] = "Yes" if row["GeneSymbol"] in benchmark_symbols else "No"
            seed_detail_rows.append(enriched)
        for gene in sorted(benchmark_symbols):
            benchmark_detail_rows.append(
                {
                    "Cancer": disease,
                    "OncoKBGene": gene,
                    "InSeedSet": "Yes" if gene in seed_symbols else "No",
                }
            )
        top100_detail_rows.extend(nonseed_ranking)

        write_tsv(
            details_dir / f"{disease}_enriched_seed_set_mapped.tsv",
            [row | {"InCancerSpecificOncoKB": "Yes" if row["GeneSymbol"] in benchmark_symbols else "No"} for row in seed_rows],
            ["Cancer", "SeedIdentifier", "SeedIdentifierBase", "GeneSymbol", "InCancerSpecificOncoKB"],
        )
        write_tsv(
            details_dir / f"{disease}_cancer_specific_oncokb_genes.tsv",
            [{"Cancer": disease, "OncoKBGene": gene, "InSeedSet": "Yes" if gene in seed_symbols else "No"} for gene in sorted(benchmark_symbols)],
            ["Cancer", "OncoKBGene", "InSeedSet"],
        )
        write_tsv(
            details_dir / f"{disease}_candidate1_top100_nonseed.tsv",
            nonseed_ranking,
            ["Cancer", "NonSeedRank", "OriginalRank", "GeneIdentifier", "GeneSymbol", "Score", "InBenchmarkExcludingSeeds"],
        )

    table_headers = [
        "Cancer",
        "Enriched seeds",
        "|Omega_c|",
        "|S_c intersect Omega_c|",
        "Seed overlap (%)",
        "Benchmark coverage (%)",
        "Matches@100 excluding seeds",
        "Recall@100 excluding seeds",
    ]
    manifest_headers = [
        "Cancer",
        "Seed file",
        "Benchmark OncoKB file",
        "Ranking file",
        "Optimization summary",
        "Alpha",
        "Beta",
        "Seed identifiers",
        "Mapped seed symbols",
        "Benchmark genes",
    ]
    write_tsv(OUTPUT_DIR / "table3_cancer_specific_seed_benchmark_overlap.tsv", table_rows, table_headers)
    write_tsv(OUTPUT_DIR / "table3_source_manifest.tsv", manifest_rows, manifest_headers)

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Table 3"
    sheet.append(table_headers)
    for row in table_rows:
        sheet.append([row[header] for header in table_headers])
    manifest_sheet = workbook.create_sheet("Source manifest")
    manifest_sheet.append(manifest_headers)
    for row in manifest_rows:
        manifest_sheet.append([row[header] for header in manifest_headers])
    seed_sheet = workbook.create_sheet("Mapped seeds")
    seed_headers = ["Cancer", "SeedIdentifier", "SeedIdentifierBase", "GeneSymbol", "InCancerSpecificOncoKB"]
    seed_sheet.append(seed_headers)
    for row in seed_detail_rows:
        seed_sheet.append([row[header] for header in seed_headers])
    benchmark_sheet = workbook.create_sheet("OncoKB gene sets")
    benchmark_headers = ["Cancer", "OncoKBGene", "InSeedSet"]
    benchmark_sheet.append(benchmark_headers)
    for row in benchmark_detail_rows:
        benchmark_sheet.append([row[header] for header in benchmark_headers])
    top100_sheet = workbook.create_sheet("Top100 non-seed")
    top100_headers = ["Cancer", "NonSeedRank", "OriginalRank", "GeneIdentifier", "GeneSymbol", "Score", "InBenchmarkExcludingSeeds"]
    top100_sheet.append(top100_headers)
    for row in top100_detail_rows:
        top100_sheet.append([row[header] for header in top100_headers])

    for ws in workbook.worksheets:
        style_sheet(ws)
    workbook.save(OUTPUT_DIR / "table3_cancer_specific_seed_benchmark_overlap.xlsx")

    print(OUTPUT_DIR / "table3_cancer_specific_seed_benchmark_overlap.xlsx")
    for row in table_rows:
        print(row)


if __name__ == "__main__":
    main()
