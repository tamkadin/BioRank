import csv
from collections import Counter
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parents[1]
DATA_DIR = ROOT / "data_set"
OUTPUT_DIR = ROOT / "output" / "report_exports" / "20260907_batch_4disease_seed_metrics"
OUTPUT_PATH = OUTPUT_DIR / "oncokb_evaluation_set_descriptive_summary.xlsx"

DISEASE_FILES = {
    "BRCA": DATA_DIR / "Onco_KB BRCA.csv",
    "COAD": DATA_DIR / "Onco_KB_COAD.csv",
    "LUAD": DATA_DIR / "Onco_KB_LUAD.csv",
    "THCA": DATA_DIR / "Onco_KB_THCA.csv",
}
SHARED_FILE = DATA_DIR / "Onco_KB.csv"


def load_genes(path):
    genes = []
    with open(path, encoding="utf-8-sig", newline="") as handle:
        reader = csv.DictReader(handle)
        for row in reader:
            gene = (row.get("Gene") or "").strip().upper()
            if gene:
                genes.append(gene)
    return genes


def style_sheet(sheet):
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    thin = Side(style="thin", color="D9E2F3")

    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for row in sheet.iter_rows():
        for cell in row:
            cell.border = Border(top=thin, bottom=thin, left=thin, right=thin)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if isinstance(cell.value, float):
                cell.number_format = "0.00%"

    sheet.freeze_panes = "A2"
    for column_index, column in enumerate(sheet.columns, 1):
        max_len = max(len(str(cell.value)) if cell.value is not None else 0 for cell in column)
        sheet.column_dimensions[get_column_letter(column_index)].width = min(max(max_len + 2, 12), 65)


def main():
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    shared_genes = load_genes(SHARED_FILE)
    shared_set = set(shared_genes)

    summary_rows = []
    gene_rows = []
    for disease, file_path in DISEASE_FILES.items():
        genes = load_genes(file_path)
        counts = Counter(genes)
        unique_genes = sorted(counts)
        overlap = sorted(set(unique_genes) & shared_set)
        disease_only = sorted(set(unique_genes) - shared_set)

        summary_rows.append(
            {
                "Disease": disease,
                "Evaluation file": str(file_path),
                "Total rows": len(genes),
                "Unique genes": len(unique_genes),
                "Duplicate genes": len(genes) - len(unique_genes),
                "Genes overlapping shared OncoKB": len(overlap),
                "Overlap within disease set (%)": len(overlap) / len(unique_genes) if unique_genes else 0,
                "Coverage of shared OncoKB (%)": len(overlap) / len(shared_set) if shared_set else 0,
                "Genes not in shared OncoKB": len(disease_only),
                "Gene list": ", ".join(unique_genes),
                "Note": "Disease-specific validation set used in the 2026-09-05 batch run.",
            }
        )
        for index, gene in enumerate(unique_genes, 1):
            gene_rows.append(
                {
                    "Disease": disease,
                    "STT": index,
                    "Gene": gene,
                    "In shared OncoKB": "Yes" if gene in shared_set else "No",
                }
            )

    workbook = Workbook()
    summary_sheet = workbook.active
    summary_sheet.title = "Summary"
    summary_headers = list(summary_rows[0])
    summary_sheet.append(summary_headers)
    for row in summary_rows:
        summary_sheet.append([row[header] for header in summary_headers])

    gene_sheet = workbook.create_sheet("Gene lists")
    gene_headers = ["Disease", "STT", "Gene", "In shared OncoKB"]
    gene_sheet.append(gene_headers)
    for row in gene_rows:
        gene_sheet.append([row[header] for header in gene_headers])

    shared_sheet = workbook.create_sheet("Shared OncoKB")
    shared_sheet.append(["Metric", "Value"])
    shared_sheet.append(["Shared OncoKB file", str(SHARED_FILE)])
    shared_sheet.append(["Total rows", len(shared_genes)])
    shared_sheet.append(["Unique genes", len(shared_set)])
    shared_sheet.append(["Duplicate genes", len(shared_genes) - len(shared_set)])
    shared_sheet.append(["Gene column", "Gene"])

    for sheet in workbook.worksheets:
        style_sheet(sheet)
        if sheet.max_column and sheet.max_row:
            sheet.auto_filter.ref = sheet.dimensions

    workbook.save(OUTPUT_PATH)
    print(OUTPUT_PATH)
    for row in summary_rows:
        print(
            row["Disease"],
            row["Unique genes"],
            row["Genes overlapping shared OncoKB"],
            f"{row['Coverage of shared OncoKB (%)']:.2%}",
            row["Gene list"],
            sep="\t",
        )


if __name__ == "__main__":
    main()
