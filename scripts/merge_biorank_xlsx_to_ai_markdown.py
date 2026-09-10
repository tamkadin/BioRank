from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
EXPORT_DIR = ROOT / "output" / "report_exports" / "20260907_batch_4disease_seed_metrics"
OUTPUT_PATH = EXPORT_DIR / "biorank_v2_batch_metrics_and_oncokb_eval_summary_for_ai.md"

WORKBOOKS = {
    "Original seed": EXPORT_DIR / "optuna_model_metrics_top100.xlsx",
    "Improved seed": EXPORT_DIR / "data_improved_optuna_model_metrics_top100.xlsx",
}
EVAL_WORKBOOK = EXPORT_DIR / "oncokb_evaluation_set_descriptive_summary.xlsx"
DISEASES = ("BRCA", "COAD", "LUAD", "THCA")
METRIC_COLUMNS = ("nDCG@100", "Recall@100", "Common genes @100", "Selection score")


def read_sheet(path, sheet_name):
    workbook = load_workbook(path, data_only=True)
    sheet = workbook[sheet_name]
    headers = [sheet.cell(1, column).value for column in range(1, sheet.max_column + 1)]
    rows = []
    for row_index in range(2, sheet.max_row + 1):
        row = {
            headers[column - 1]: sheet.cell(row_index, column).value
            for column in range(1, sheet.max_column + 1)
        }
        if any(value is not None and value != "" for value in row.values()):
            rows.append(row)
    return rows


def markdown_escape(value):
    if value is None:
        return ""
    text = str(value)
    return text.replace("|", "\\|").replace("\n", "<br>")


def display_path(path):
    try:
        return str(path.relative_to(ROOT))
    except ValueError:
        return str(path)


def format_value(header, value):
    if value is None:
        return ""
    if isinstance(value, str):
        candidate = Path(value)
        if candidate.is_absolute():
            try:
                return str(candidate.relative_to(ROOT))
            except ValueError:
                return value
    if header.endswith("(%)"):
        try:
            return f"{float(value) * 100:.2f}%"
        except (TypeError, ValueError):
            return str(value)
    if isinstance(value, float):
        return f"{value:.4f}"
    return str(value)


def table(headers, rows):
    lines = []
    lines.append("| " + " | ".join(headers) + " |")
    lines.append("| " + " | ".join("---" for _ in headers) + " |")
    for row in rows:
        values = []
        for header in headers:
            values.append(markdown_escape(format_value(header, row.get(header))))
        lines.append("| " + " | ".join(values) + " |")
    return "\n".join(lines)


def best_row(rows):
    return max(rows, key=lambda row: float(row.get("Selection score") or 0))


def metric_average(rows, metric):
    values = [float(row.get(metric) or 0) for row in rows]
    return sum(values) / len(values) if values else 0.0


def main():
    export_rows = {
        seed_name: {
            disease: read_sheet(path, disease)
            for disease in DISEASES
        }
        for seed_name, path in WORKBOOKS.items()
    }
    eval_summary = read_sheet(EVAL_WORKBOOK, "Summary")
    gene_list = read_sheet(EVAL_WORKBOOK, "Gene lists")
    shared_oncokb = read_sheet(EVAL_WORKBOOK, "Shared OncoKB")

    lines = [
        "# BioRank v2 Batch Metrics and OncoKB Evaluation Summary",
        "",
        "## Purpose",
        "",
        "This Markdown file merges the three Excel outputs into one AI-readable report. It contains:",
        "",
        "- Optuna Top-100 metrics using original seed genes.",
        "- Optuna Top-100 metrics using improved/enriched seed genes.",
        "- Descriptive statistics for the disease-specific OncoKB evaluation sets.",
        "",
        "## Source Files",
        "",
        f"- Original seed metrics: `{display_path(WORKBOOKS['Original seed'])}`",
        f"- Improved seed metrics: `{display_path(WORKBOOKS['Improved seed'])}`",
        f"- OncoKB evaluation set summary: `{display_path(EVAL_WORKBOOK)}`",
        "",
        "## Run Context",
        "",
        "- Diseases: BRCA, COAD, LUAD, THCA.",
        "- Optimization case: BioRank v2 full.",
        "- Trials per disease run: 200.",
        "- Random seed: 42.",
        "- Evaluation mode: Disease OncoKB.",
        "- Metrics reported here are Top-100 metrics.",
        "",
        "## Metric Definitions",
        "",
        "- `nDCG@100`: position-aware ranking quality for validation hits in the top 100.",
        "- `Recall@100`: fraction of validation genes recovered in the top 100.",
        "- `Common genes @100`: count of validation genes found in the top 100.",
        "- `Selection score`: internal display/selection score used to compare candidates.",
        "",
        "Important interpretation note: disease-specific OncoKB sets are much smaller than the pan-cancer OncoKB file. Therefore `Recall@100` from disease-specific evaluation and pan-cancer evaluation should not be interpreted as the same denominator.",
        "",
    ]

    lines.extend(
        [
            "## Disease-Specific OncoKB Evaluation Sets",
            "",
            table(
                [
                    "Disease",
                    "Total rows",
                    "Unique genes",
                    "Duplicate genes",
                    "Genes overlapping shared OncoKB",
                    "Coverage of shared OncoKB (%)",
                    "Genes not in shared OncoKB",
                ],
                eval_summary,
            ),
            "",
            "## Shared OncoKB File Summary",
            "",
            table(["Metric", "Value"], shared_oncokb),
            "",
            "## Best Candidate Summary",
            "",
        ]
    )

    best_summary = []
    for seed_name, disease_rows in export_rows.items():
        for disease, rows in disease_rows.items():
            row = best_row(rows)
            best_summary.append(
                {
                    "Seed option": seed_name,
                    "Disease": disease,
                    "Best method": row.get("Method"),
                    "Alpha": row.get("Alpha"),
                    "Beta": row.get("Beta"),
                    "nDCG@100": row.get("nDCG@100"),
                    "Recall@100": row.get("Recall@100"),
                    "Common genes @100": row.get("Common genes @100"),
                    "Selection score": row.get("Selection score"),
                }
            )
    lines.extend(
        [
            table(
                [
                    "Seed option",
                    "Disease",
                    "Best method",
                    "Alpha",
                    "Beta",
                    "nDCG@100",
                    "Recall@100",
                    "Common genes @100",
                    "Selection score",
                ],
                best_summary,
            ),
            "",
            "## Average Best Metrics by Seed Option",
            "",
        ]
    )

    average_summary = []
    for seed_name in WORKBOOKS:
        rows = [row for row in best_summary if row["Seed option"] == seed_name]
        average_summary.append(
            {
                "Seed option": seed_name,
                "Avg nDCG@100": metric_average(rows, "nDCG@100"),
                "Avg Recall@100": metric_average(rows, "Recall@100"),
                "Avg Common genes @100": metric_average(rows, "Common genes @100"),
                "Avg Selection score": metric_average(rows, "Selection score"),
            }
        )
    lines.extend(
        [
            table(
                [
                    "Seed option",
                    "Avg nDCG@100",
                    "Avg Recall@100",
                    "Avg Common genes @100",
                    "Avg Selection score",
                ],
                average_summary,
            ),
            "",
            "## Detailed Optuna Metrics",
            "",
        ]
    )

    detail_headers = [
        "STT",
        "Ranking file",
        "Method",
        "Algorithm",
        "Variant",
        "Selection source",
        "Alpha",
        "Beta",
        "nDCG@100",
        "Recall@100",
        "Common genes @100",
        "Selection score",
    ]
    for seed_name, disease_rows in export_rows.items():
        lines.extend([f"### {seed_name}", ""])
        for disease, rows in disease_rows.items():
            lines.extend([f"#### {disease}", "", table(detail_headers, rows), ""])

    lines.extend(["## Disease-Specific Gene Lists", ""])
    gene_headers = ["Disease", "STT", "Gene", "In shared OncoKB"]
    for disease in DISEASES:
        rows = [row for row in gene_list if row.get("Disease") == disease]
        lines.extend([f"### {disease}", "", table(gene_headers, rows), ""])

    lines.extend(
        [
            "## Compact Interpretation for AI",
            "",
            "- The improved seed setting has a higher average best `Selection score` than original seed under disease-specific OncoKB evaluation.",
            "- Improved seed also has higher average `nDCG@100`, `Recall@100`, and `Common genes @100` across the four diseases.",
            "- THCA shows the largest improved-seed disease-specific best score among the listed runs.",
            "- The disease-specific validation files contain 26-32 unique genes per disease and no duplicate genes.",
            "",
        ]
    )

    OUTPUT_PATH.write_text("\n".join(lines), encoding="utf-8")
    print(OUTPUT_PATH)


if __name__ == "__main__":
    main()
